from docx import Document

def create_word_doc():
    doc = Document()
    doc.add_heading('List of 47 Continuous Features (Ranked by Priority)', 0)

    features_by_tier = {
        "Tier 1: Critical Real-Time Indicators": [
            "1. shock_index", "2. respiration", "3. systemicsystolic", 
            "4. systemicdiastolic", "5. systemicmean", "6. heartrate", 
            "7. map_calculated", "8. pulse_pressure"
        ],
        "Tier 2: Organ Damage & Infection Markers": [
            "9. Lactate", "10. Creatinine", "11. Bilirubin", 
            "12. WBC", "13. Platelets", "14. temperature"
        ],
        "Tier 3: The 6-Hour Trajectories": [
            "15. heartrate_rate_of_change", "16. systemicsystolic_rate_of_change", 
            "17. systemicdiastolic_rate_of_change", "18. systemicmean_rate_of_change", 
            "19. respiration_rate_of_change", "20. temperature_rate_of_change", 
            "21. heartrate_rolling_std", "22. systemicsystolic_rolling_std", 
            "23. systemicdiastolic_rolling_std", "24. systemicmean_rolling_std", 
            "25. respiration_rolling_std", "26. temperature_rolling_std", 
            "27. heartrate_rolling_mean", "28. systemicsystolic_rolling_mean", 
            "29. systemicdiastolic_rolling_mean", "30. systemicmean_rolling_mean", 
            "31. respiration_rolling_mean", "32. temperature_rolling_mean"
        ],
        "Tier 4: Routine Labs & Resp Support": [
            "33. resp_fio2", "34. resp_fio2_(%)", "35. resp_peep", 
            "36. Hemoglobin", "37. Potassium", "38. Sodium", 
            "39. Glucose", "40. resp_peep/cpap", "41. resp_ps_above_peep", 
            "42. resp_set_fraction_of_inspired_oxygen_(fio2)", 
            "43. resp_unable_to_obtain_peepi_and_vtrap"
        ],
        "Tier 5: Static Measurements": [
            "44. age", "45. bmi", "46. admissionweight", "47. admissionheight"
        ]
    }

    for tier, features in features_by_tier.items():
        doc.add_heading(tier, level=1)
        for feature in features:
            doc.add_paragraph(feature, style='List Number')

    doc.save('47_columns_list.docx')
    print("Created 47_columns_list.docx successfully.")

if __name__ == "__main__":
    create_word_doc()
